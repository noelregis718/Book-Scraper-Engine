import asyncio
import pandas as pd
import re
import os
import urllib.request
import urllib.parse
import json
from playwright.async_api import async_playwright

EXCEL_FILE = r"e:\Internship\PocketFM\book_details_from_email_rechecked_filled.xlsx"
START_ROW = 0
TARGET_ROWS = 5
CONCURRENCY = 2
BATCH_SIZE = 5

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'application/json, text/javascript, */*; q=0.01',
}

def get_autocomplete_book_url(query):
    import time
    api_url = f"https://www.goodreads.com/book/auto_complete?format=json&q={urllib.parse.quote_plus(query)}"
    for attempt in range(3):
        try:
            req = urllib.request.Request(api_url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=15) as response:
                data = json.loads(response.read().decode("utf-8", errors="ignore"))
                if data and len(data) > 0:
                    book_path = data[0].get('bookUrl', '')
                    if book_path:
                        return "https://www.goodreads.com" + book_path
                return None
        except Exception as e:
            time.sleep(1)
    return None

async def process_row(index, row, df, context, sem):
    async with sem:
        series_name = str(row.get("Series Name", "")).strip()
        book_name = str(row.get("Book Name", "")).strip()
        author_name = str(row.get("Author Name", "")).strip()
        
        # Determine best search term
        if series_name and series_name.lower() != 'nan':
            base_query = series_name
        elif book_name and book_name.lower() != 'nan':
            base_query = book_name
        else:
            print(f"[{index}] No Series Name or Book Name found. Skipping.")
            return
            
        if ":" in base_query:
            base_query = base_query.split(":")[0].strip()
            
        if not author_name or author_name.lower() in ['nan', 'none', '']:
            search_query = base_query
        else:
            search_query = f"{base_query} {author_name}"
            
        print(f"\n--- Processing Row {index + 2} (Index {index}) ---")
        print(f"[{index}] Querying: '{search_query}'")

        # 1. Bypass WAF using Autocomplete API
        book_url = await asyncio.to_thread(get_autocomplete_book_url, search_query)
        if not book_url:
            print(f"[{index}] Book not found in Autocomplete API. Skipping for now.")
            return
            
        print(f"[{index}] Found Book URL: {book_url}")
        
        # 2. Load Book Page using Playwright
        page = await context.new_page()
        try:
            await page.goto(book_url, wait_until="domcontentloaded", timeout=45000)
            await asyncio.sleep(2)
                
            # EXTRACT RATING
            avg_rating = 0.0
            rating_el = await page.query_selector('div.RatingStatistics__rating')
            if rating_el:
                try: avg_rating = float((await rating_el.inner_text()).strip())
                except: pass
                
            df.at[index, "Book1_Rating"] = avg_rating
            
            # Find Series Link
            series_tag = await page.query_selector('h3.Text__title3 a[href*="/series/"], [data-testid="series"] a, div.BookPageTitleSection__title a[href*="/series/"], a.infoBoxRowItem[href*="/series/"]')
            if not series_tag:
                print(f"[{index}] No Series link found. Processing as standalone.")
                df.at[index, "GoodReads_Series_URL"] = book_url
                
                # Standalone page extraction
                pages_el = await page.query_selector('[data-testid="pagesFormat"]')
                if pages_el:
                    pages_text = await pages_el.inner_text()
                    p_match = re.search(r'(\d+)\s*pages', pages_text, re.IGNORECASE)
                    if p_match:
                        pages = int(p_match.group(1))
                        df.at[index, "Num_Primary_Books_in_Series"] = 1
                        df.at[index, "Total_Page_Count_of_Primary_Books"] = pages
                        df.at[index, "Approx Length (Hrs)"] = round((pages * 250) / 10000, 1)
                        print(f"[{index}] Standalone Pages: {pages}")
                return
                
            # Extract Series URL and navigate
            series_url = await series_tag.evaluate("el => el.href")
            print(f"[{index}] Found Series Link: {series_url}")
            df.at[index, "GoodReads_Series_URL"] = series_url
            
            await page.goto(series_url, wait_until="domcontentloaded", timeout=45000)
            await asyncio.sleep(1)
            
            # Parse Series List
            num_primary_books = 0
            total_page_count = 0
            
            book_items = await page.query_selector_all('.listWithDividers__item, .seriesWork')
            
            for item in book_items:
                item_text = await item.inner_text()
                match = re.search(r'Book\s+([0-9a-zA-Z\.\-]+)', item_text, re.IGNORECASE)
                if match:
                    book_num = match.group(1)
                    if book_num.isdigit():
                        num_primary_books += 1
                        
                        b_link = await item.query_selector('a.bookTitle, a[href*="/book/show"]')
                        if b_link:
                            b_url = await b_link.evaluate("el => el.href")
                            b_page = await context.new_page()
                            try:
                                await b_page.goto(b_url, wait_until="domcontentloaded", timeout=30000)
                                await asyncio.sleep(1)
                                
                                extracted_pages = 0
                                # JSON-LD
                                try:
                                    ld_el = await b_page.query_selector('script[type="application/ld+json"]')
                                    if ld_el:
                                        data = json.loads(await ld_el.inner_text())
                                        if isinstance(data, list): data = data[0]
                                        if 'numberOfPages' in data:
                                            extracted_pages = int(data['numberOfPages'])
                                except: pass
                                
                                # Fallback regex
                                if not extracted_pages:
                                    try:
                                        content = await b_page.content()
                                        matches = re.findall(r'(\d+)\s*pages', content, re.IGNORECASE)
                                        if matches:
                                            extracted_pages = int(matches[0])
                                    except: pass

                                if extracted_pages:
                                    total_page_count += extracted_pages
                                    print(f"[{index}]   - Book {book_num} | {extracted_pages} pages")
                            except Exception as e:
                                print(f"[{index}]   - Failed to scrape Book {book_num}: {e}")
                            finally:
                                await b_page.close()
                                    
            if num_primary_books == 0:
                num_primary_books = max(1, len(book_items))
                
            print(f"[{index}] Complete. Primary Books: {num_primary_books} | Total Pages: {total_page_count}")
            df.at[index, "Num_Primary_Books_in_Series"] = num_primary_books
            df.at[index, "Total_Page_Count_of_Primary_Books"] = total_page_count
            if total_page_count > 0:
                df.at[index, "Approx Length (Hrs)"] = round((total_page_count * 250) / 10000, 1)

        except Exception as e:
            print(f"[{index}] Error: {e}")
        finally:
            await page.close()

async def run_scraper():
    print(f"Loading {EXCEL_FILE}...")
    try:
        # Load with header at row 1 (index 1)
        df = pd.read_excel(EXCEL_FILE, header=1)
    except Exception as e:
        print(f"Excel load error: {e}")
        return

    # Ensure all target columns exist (K onwards)
    new_cols = [
        "GoodReads_Series_URL",
        "Book1_Rating",
        "Num_Primary_Books_in_Series",
        "Total_Page_Count_of_Primary_Books",
        "Approx Length (Hrs)"
    ]
    for col in new_cols:
        if col not in df.columns:
            df[col] = None

    print(f"Running custom scraper on FIRST {TARGET_ROWS} ROWS...")

    async with async_playwright() as p:
        context = await p.chromium.launch_persistent_context(
            user_data_dir=os.path.join(r"e:\Internship\PocketFM", "browser_profile_main"),
            headless=False,
            args=['--disable-blink-features=AutomationControlled'],
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36",
            viewport={'width': 1280, 'height': 800}
        )
        
        sem = asyncio.Semaphore(CONCURRENCY)
        tasks = []
        for index in range(START_ROW, min(START_ROW + TARGET_ROWS, len(df))):
            row = df.iloc[index]
            tasks.append(process_row(index, row, df, context, sem))
            
        await asyncio.gather(*tasks)
        await context.close()
        
    try:
        # We need to preserve the first row (the main title). 
        # Reading with header=1 strips the row 0 title.
        # Let's save it back with the header, it will override the multi-level header but that's standard for pandas.
        # If preserving the exact layout is needed, openpyxl is better.
        # For simplicity and correctness with pandas:
        # We will load the original file with openpyxl, and just write the 5 columns back.
        import openpyxl
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        
        # Find starting column (K is column 11)
        start_col = 11
        
        # Write Headers on row 2
        for c_idx, col_name in enumerate(new_cols):
            ws.cell(row=2, column=start_col + c_idx, value=col_name)
            
        # Write Data
        for index in range(START_ROW, min(START_ROW + TARGET_ROWS, len(df))):
            row_idx = index + 3  # index 0 in dataframe is row 3 in excel (since header is row 2)
            ws.cell(row=row_idx, column=start_col, value=df.at[index, "GoodReads_Series_URL"])
            ws.cell(row=row_idx, column=start_col + 1, value=df.at[index, "Book1_Rating"])
            ws.cell(row=row_idx, column=start_col + 2, value=df.at[index, "Num_Primary_Books_in_Series"])
            ws.cell(row=row_idx, column=start_col + 3, value=df.at[index, "Total_Page_Count_of_Primary_Books"])
            ws.cell(row=row_idx, column=start_col + 4, value=df.at[index, "Approx Length (Hrs)"])
            
        wb.save(EXCEL_FILE)
        print("\nSuccessfully updated Excel file with new columns!")
        
    except Exception as e:
        print(f"Failed to save to Excel: {e}")

if __name__ == "__main__":
    asyncio.run(run_scraper())
