import openpyxl
import requests
from bs4 import BeautifulSoup
import re
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib3.util.retry import Retry
from requests.adapters import HTTPAdapter

CACHE_FILE = r'C:\Users\noelr\.gemini\antigravity-ide\brain\54260b81-f2ff-4d78-b4cd-0f6fd4f86372\scratch\picks_gr_cache.json'
TARGET_FILES = [
    'Romantasy - Picks for Review.xlsx',
    'Romantasy - Subjective Reviews - 9.9.26 - New Sheet.xlsx'
]

# Load existing cache
cache = {}
if os.path.exists(CACHE_FILE):
    try:
        cache = json.load(open(CACHE_FILE, encoding='utf-8'))
        print(f"Loaded {len(cache)} entries from cache.")
    except:
        pass

def create_session():
    session = requests.Session()
    retries = Retry(total=3, backoff_factor=1.2, status_forcelist=[429, 500, 502, 503, 504])
    session.mount('https://', HTTPAdapter(max_retries=retries))
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
        'Accept-Language': 'en-US,en;q=0.9',
    })
    return session

def fetch_gr_series_details(series_url):
    if series_url in cache and cache[series_url].get('status') == 'SUCCESS':
        return series_url, cache[series_url]

    session = create_session()
    details = {
        'status': 'FAILED',
        'title': None,
        'author': None,
        'primary_books_count': None,
        'book1_link': None,
        'book1_title': None,
        'book1_rating': None,
        'book1_pages': None,
        'book1_ratings_count': None,
        'error': None
    }
    
    if '/book/show/' in series_url:
        try:
            r = session.get(series_url, timeout=15)
            if r.status_code == 200:
                soup = BeautifulSoup(r.text, 'html.parser')
                for s in soup.select('script[type="application/ld+json"]'):
                    try:
                        data = json.loads(s.string)
                        if data.get('@type') == 'Book':
                            details['book1_title'] = data.get('name')
                            details['title'] = data.get('name')
                            details['book1_link'] = series_url
                            details['book1_pages'] = data.get('numberOfPages')
                            details['book1_rating'] = data.get('aggregateRating', {}).get('ratingValue')
                            details['book1_ratings_count'] = data.get('aggregateRating', {}).get('ratingCount')
                            details['primary_books_count'] = 1
                            if data.get('author'):
                                a_info = data['author']
                                if isinstance(a_info, list) and len(a_info) > 0:
                                    details['author'] = a_info[0].get('name')
                                elif isinstance(a_info, dict):
                                    details['author'] = a_info.get('name')
                            details['status'] = 'SUCCESS'
                            break
                    except:
                        pass
        except Exception as e:
            details['error'] = str(e)
        return series_url, details

    try:
        r = session.get(series_url, timeout=15)
        if r.status_code == 200:
            soup = BeautifulSoup(r.text, 'html.parser')
            h1 = soup.select_one('h1')
            s_name = h1.get_text(strip=True) if h1 else ''
            if s_name.lower().endswith(" series"):
                s_name = s_name[:-7].strip()
            details['title'] = s_name
            
            for a in soup.select('a[href*="/author/show/"]'):
                name = a.get_text(strip=True)
                if name and not name.startswith('http') and len(name) > 2:
                    details['author'] = name
                    break
            if not details['author'] and soup.title:
                m_auth = re.search(r'by\s+([^|•\n\r]+)', soup.title.string or '')
                if m_auth:
                    details['author'] = m_auth.group(1).strip()

            header_text = soup.select_one('.responsiveSeriesHeader, .seriesHeader')
            htext = header_text.get_text(separator=' ', strip=True) if header_text else soup.get_text()
            m_works = re.search(r'(\d+)\s+primary works?', htext)
            if m_works:
                details['primary_books_count'] = int(m_works.group(1))

            b1_el = soup.select_one('.listWithDividers__item a.gr-h3, .listWithDividers__item a[href*="/book/show/"]')
            if b1_el and b1_el.get('href'):
                b1_url = "https://www.goodreads.com" + b1_el['href'].split('?')[0]
                details['book1_link'] = b1_url
                
                time.sleep(0.15)
                r_b1 = session.get(b1_url, timeout=15)
                if r_b1.status_code == 200:
                    soup_b1 = BeautifulSoup(r_b1.text, 'html.parser')
                    for s in soup_b1.select('script[type="application/ld+json"]'):
                        try:
                            data = json.loads(s.string)
                            if data.get('@type') == 'Book':
                                details['book1_title'] = data.get('name')
                                details['book1_pages'] = data.get('numberOfPages')
                                details['book1_rating'] = data.get('aggregateRating', {}).get('ratingValue')
                                details['book1_ratings_count'] = data.get('aggregateRating', {}).get('ratingCount')
                                if not details['author'] and data.get('author'):
                                    a_info = data['author']
                                    if isinstance(a_info, list) and len(a_info) > 0:
                                        details['author'] = a_info[0].get('name')
                                    elif isinstance(a_info, dict):
                                        details['author'] = a_info.get('name')
                                break
                        except:
                            pass
            details['status'] = 'SUCCESS'
    except Exception as e:
        details['error'] = str(e)
    return series_url, details

# 1. Collect priority error rows
wb_ref = openpyxl.load_workbook('Romantasy - Picks for Review.xlsx', data_only=True)
ws_ref = wb_ref['Picks for Review (Sep 26)']

priority_urls = set()
for r in range(2, ws_ref.max_row + 1):
    gr_url = str(ws_ref.cell(r, 4).value or '').strip()
    gr_ratings = ws_ref.cell(r, 5).value
    b1_rating = ws_ref.cell(r, 6).value
    page_count = ws_ref.cell(r, 8).value
    
    needs = False
    # A. Ratings == Page Count
    if gr_ratings and page_count and str(gr_ratings).strip() == str(page_count).strip():
        needs = True
    # B. Unrealistic or missing page count
    try:
        if page_count is None or str(page_count).strip() == '' or float(str(page_count).replace(',', '')) > 15000 or float(str(page_count).replace(',', '')) < 50:
            needs = True
    except:
        needs = True
    # C. Missing or invalid B1 rating
    if b1_rating is None or str(b1_rating).strip() == '':
        needs = True
    else:
        try:
            v = float(str(b1_rating).replace(',', ''))
            if v > 5.0 or v < 1.0: needs = True
        except:
            needs = True
            
    if needs and gr_url.startswith('http'):
        if gr_url not in cache:
            priority_urls.add(gr_url)

print(f"Total Priority Error URLs needing live scrape: {len(priority_urls)}")

# Scrape priority URLs with 5 workers
if priority_urls:
    print(f"Scraping {len(priority_urls)} priority Goodreads URLs...")
    c = 0
    with ThreadPoolExecutor(max_workers=5) as ex:
        futs = {ex.submit(fetch_gr_series_details, u): u for u in priority_urls}
        for f in as_completed(futs):
            u, res = f.result()
            cache[u] = res
            c += 1
            if c % 10 == 0 or c == len(priority_urls):
                print(f"  Scraped {c}/{len(priority_urls)} priority URLs...")
                json.dump(cache, open(CACHE_FILE, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
            time.sleep(0.1)
    json.dump(cache, open(CACHE_FILE, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print("Priority scraping complete!")

# 2. Apply all fixes across both files
fixes_summary = {
    'author_format': 0,
    'series_title_format': 0,
    'author_mismatch': 0,
    'dirty_publisher': 0,
    'ratings_in_page_count_fixed': 0,
    'unrealistic_page_count_fixed': 0,
    'b1_rating_fixed': 0,
    'b1_link_fixed': 0
}

# Known author corrections from misalignment
KNOWN_AUTHOR_CORRECTIONS = {
    'https://www.goodreads.com/series/206510-green-creek': 'T.J. Klune',
    'https://www.goodreads.com/series/319224-this-woven-kingdom': 'Tahereh Mafi',
    'https://www.goodreads.com/series/53144-angel': 'Angela Crocker',
}

for fn in TARGET_FILES:
    print(f"\nProcessing {fn}...")
    wb = openpyxl.load_workbook(fn)
    ws = wb['Picks for Review (Sep 26)']
    
    for r in range(2, min(1231, ws.max_row + 1)):
        s_name = str(ws.cell(r, 1).value or '').strip()
        author = str(ws.cell(r, 2).value or '').strip()
        gr_url = str(ws.cell(r, 4).value or '').strip()
        gr_ratings = ws.cell(r, 5).value
        b1_rating = ws.cell(r, 6).value
        primary_books = ws.cell(r, 7).value
        page_count = ws.cell(r, 8).value
        imprint_book = str(ws.cell(r, 9).value or '').strip()
        imprint_audio = str(ws.cell(r, 11).value or '').strip()
        b1_link = str(ws.cell(r, 16).value or '').strip()
        
        # 1. Author formatting: clean 'by '
        clean_author = re.sub(r'^(by\s+|author:\s*|translated\s+by\s*)', '', author, flags=re.IGNORECASE).strip()
        clean_author = re.sub(r'^(by\s+)', '', clean_author, flags=re.IGNORECASE).strip()
        if clean_author != author:
            ws.cell(r, 2).value = clean_author
            author = clean_author
            fixes_summary['author_format'] += 1
            
        # 2. Series title repeats: clean '... by Author by Author'
        clean_sname = re.sub(r'\s+by\s+' + re.escape(author) + r'(\s+by\s+' + re.escape(author) + r')?$', '', s_name, flags=re.IGNORECASE).strip()
        clean_sname = re.sub(r'^(Rogues of the Crosslands Book \d+ of \d+:\s*)', '', clean_sname).strip()
        clean_sname = re.sub(r'^(Sins of the Zodiac Book \d+ of \d+:\s*)', '', clean_sname).strip()
        clean_sname = re.sub(r'^(The Anna Albertini Files Book \d+ of \d+:\s*)', '', clean_sname).strip()
        clean_sname = re.sub(r'^(Briarborn Chronicles Book \d+ of \d+:\s*)', '', clean_sname).strip()
        if clean_sname != s_name and len(clean_sname) > 2:
            ws.cell(r, 1).value = clean_sname
            fixes_summary['series_title_format'] += 1

        # 3. Dirty Publisher Imprint
        if imprint_book.startswith(')') or 'format:' in imprint_book.lower():
            if not imprint_audio or imprint_audio.strip() == '':
                ws.cell(r, 11).value = 'Audible Audiobook'
            ws.cell(r, 9).value = 'Indie / Self-Published'
            fixes_summary['dirty_publisher'] += 1

        # 4. Scraped data fixes
        c_data = cache.get(gr_url, {})
        if gr_url in KNOWN_AUTHOR_CORRECTIONS:
            ws.cell(r, 2).value = KNOWN_AUTHOR_CORRECTIONS[gr_url]
            fixes_summary['author_mismatch'] += 1
        elif c_data.get('status') == 'SUCCESS' and c_data.get('author'):
            c_auth = c_data.get('author')
            if author in ['Kelly Bennett', 'Tiffany W. Killoren', 'by CB Samet', 'Unknown']:
                ws.cell(r, 2).value = c_auth
                fixes_summary['author_mismatch'] += 1
                
        if c_data.get('status') == 'SUCCESS':
            # Book 1 Link
            if c_data.get('book1_link') and (not b1_link or not b1_link.startswith('http')):
                ws.cell(r, 16).value = c_data.get('book1_link')
                fixes_summary['b1_link_fixed'] += 1

            # Book 1 Rating
            is_b1_bad = False
            if b1_rating is None or str(b1_rating).strip() == '':
                is_b1_bad = True
            else:
                try:
                    val = float(str(b1_rating).replace(',', ''))
                    if val > 5.0 or val < 1.0: is_b1_bad = True
                except:
                    is_b1_bad = True
            if is_b1_bad and c_data.get('book1_rating'):
                ws.cell(r, 6).value = float(c_data.get('book1_rating'))
                fixes_summary['b1_rating_fixed'] += 1

            # Page Count Fixes
            is_pc_ratings = (gr_ratings and page_count and str(gr_ratings).strip() == str(page_count).strip())
            is_pc_unrealistic = False
            try:
                if page_count is None or str(page_count).strip() == '':
                    is_pc_unrealistic = True
                else:
                    pval = float(str(page_count).replace(',', ''))
                    if pval > 15000 or pval < 50: is_pc_unrealistic = True
            except:
                is_pc_unrealistic = True
                
            if is_pc_ratings or is_pc_unrealistic:
                c_pages = c_data.get('book1_pages')
                pworks = c_data.get('primary_books_count') or ws.cell(r, 7).value or 1
                try:
                    pworks_int = int(pworks)
                except:
                    pworks_int = 1
                if c_pages:
                    est_total = int(c_pages) * pworks_int
                    ws.cell(r, 8).value = est_total
                    if is_pc_ratings:
                        fixes_summary['ratings_in_page_count_fixed'] += 1
                    else:
                        fixes_summary['unrealistic_page_count_fixed'] += 1

    wb.save(fn)
    print(f"Saved {fn} successfully!")

print("\n=== SANITY FIXES SUMMARY ===")
for k, v in fixes_summary.items():
    print(f"  {k}: {v}")
