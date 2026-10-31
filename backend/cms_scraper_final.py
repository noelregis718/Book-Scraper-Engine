import pandas as pd
from playwright.sync_api import sync_playwright
from bs4 import BeautifulSoup
import re
import time
import os
from dotenv import load_dotenv

def extract_metrics(html):
    soup = BeautifulSoup(html, 'html.parser')
    metrics = {
        'Show Name': '',
        'Plays': '0',
        'Duration': '0',
        'Ratings': '0',
        'Reviews': '0',
        'Comments': '0'
    }
    
    # 1. Extract Show Name
    title_tag = soup.find('p', class_='sc-imWYAI')
    if not title_tag:
        title_tag = soup.find(['h1', 'h2'])
    if title_tag: 
        metrics['Show Name'] = title_tag.get_text(strip=True)

    def clean_val(v):
        v = str(v).strip()
        if not v or "no " in v.lower() or v.lower() == 'nan':
            return "0"
        return v

    def get_metric(label_regex):
        label_divs = soup.find_all(['div', 'span', 'p'], string=re.compile(label_regex, re.I))
        for div in label_divs:
            # Check for stat label class or card structure
            is_stat_label = 'sc-fmNzgT' in div.get('class', [])
            parent = div.parent
            if parent and (is_stat_label or len(parent.find_all(recursive=False)) == 2):
                for child in parent.find_all(recursive=False):
                    if child != div and child.get_text(strip=True):
                        return child.get_text(strip=True), div.get_text(strip=True)
                val_div = parent.find('div')
                if val_div and val_div != div and val_div.get_text(strip=True):
                    return val_div.get_text(strip=True), div.get_text(strip=True)
        return '', ''

    plays, _ = get_metric(r'^\s*Plays\s*$')
    metrics['Plays'] = clean_val(plays)

    durations, _ = get_metric(r'^\s*Hours\s*$')
    metrics['Duration'] = clean_val(durations)

    comments, _ = get_metric(r'^\s*Comments\s*$')
    metrics['Comments'] = clean_val(comments)

    ratings, reviews_label = get_metric(r'^\s*(\d+\s+)?Reviews\s*$')
    metrics['Ratings'] = clean_val(ratings)
    if reviews_label:
        rev_num = re.sub(r'[^\d.]', '', reviews_label)
        metrics['Reviews'] = clean_val(rev_num)

    return metrics


def get_tracker_path():
    root = "e:/Internship/PocketFM"
    preferred = os.path.join(root, "US_Licensing_Lifecycle_Tracker.xlsx")
    if os.path.exists(preferred):
        return preferred
    candidates = [
        f for f in os.listdir(root)
        if f.startswith("US_Licensing_Lifecycle_Tracker") and f.endswith(".xlsx") and not f.startswith("~$") and "backup" not in f.lower()
    ]
    if candidates:
        candidates.sort(key=lambda x: [int(s) if s.isdigit() else s for s in re.findall(r'\d+|\D+', x)], reverse=True)
        return os.path.join(root, candidates[0])
    return preferred


def main():
    load_dotenv()
    EMAIL = os.getenv("CMS_EMAIL")
    PASSWORD = os.getenv("CMS_PASSWORD")

    excel_path = get_tracker_path()
    sheet_name = "Argus Show Data"
    
    print("=" * 60)
    print(f" TARGET WORKBOOK: {excel_path}")
    print(f" TARGET TAB:      {sheet_name} (ONLY)")
    print("=" * 60)
    
    auth_file = "e:/Internship/PocketFM/scratch/auth.json"
    os.makedirs("e:/Internship/PocketFM/scratch", exist_ok=True)
    
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        
        if os.path.exists(auth_file):
            context = browser.new_context(storage_state=auth_file)
            print("Loaded saved login session from scratch/auth.json")
        else:
            context = browser.new_context()
            
        page = context.new_page()
        
        # Verify if session is currently valid
        session_valid = False
        if os.path.exists(auth_file):
            print("Verifying session validity on PocketFM CMS...")
            try:
                page.goto("https://cms.pocketfm.com/", wait_until="domcontentloaded", timeout=20000)
                page.wait_for_timeout(3000)
                body_text = page.inner_text("body").lower()
                has_login_prompt = (
                    "login with jumpcloud" in body_text 
                    or "please login" in body_text 
                    or "logout" in page.url.lower()
                    or "jumpcloud.com" in page.url.lower()
                    or page.locator('button:has-text("JumpCloud"), a:has-text("JumpCloud")').count() > 0
                )
                if not has_login_prompt:
                    session_valid = True
                    print("--> Saved login session is active and authenticated!")
                else:
                    print("--> Saved session has expired (Login screen detected). Opening JumpCloud SSO...")
            except Exception as e:
                print(f"--> Session verification notice: {e}")
        
        if not session_valid:
            print("\nOpening PocketFM CMS for SSO Login...")
            page.goto("https://cms.pocketfm.com/")
            print("\nAutomating SSO Login...")
            try:
                login_button = page.locator('button:has-text("JumpCloud"), a:has-text("JumpCloud"), button:has-text("Login"), a:has-text("Login"), button:has-text("SSO")')
                if login_button.first.is_visible(timeout=5000):
                    login_button.first.click()
                
                page.wait_for_url("**/sso.jumpcloud.com/**", timeout=15000)
                
                email_input = page.locator('input[type="email"], input[name="email"], input[name="username"], input[placeholder*="email" i]')
                email_input.first.wait_for(state="visible", timeout=10000)
                if EMAIL:
                    email_input.first.fill(EMAIL)
                
                next_btn = page.locator('button:has-text("Next"), button:has-text("Continue")')
                if next_btn.first.is_visible(timeout=3000):
                    next_btn.first.click()
                    page.wait_for_timeout(1000)
                
                pwd_input = page.locator('input[type="password"], input[name="password"]')
                pwd_input.first.wait_for(state="visible", timeout=5000)
                if PASSWORD:
                    pwd_input.first.fill(PASSWORD)
                
                submit_btn = page.locator('button[type="submit"], button:has-text("Log In"), button:has-text("Sign In"), a:has-text("Log In")')
                if submit_btn.first.is_visible(timeout=3000):
                    submit_btn.first.click()
                
            except Exception as e:
                print(f"\nCould not auto-fill all login fields: {e}")
                
            print("\n" + "="*60)
            print(" ACTION REQUIRED: PLEASE ENTER AUTH CODE / 2FA")
            print("="*60)
            input("The browser window is open. Please complete any 2FA / Auth Code in the browser.\nWait until the CMS dashboard renders and you no longer see JumpCloud redirects.\nOnce the dashboard is loaded, press ENTER here to start scraping: ")
            
            time.sleep(3)
            context.storage_state(path=auth_file)
            print("Login session saved successfully to scratch/auth.json!")
            
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, data_only=False)
        ws = wb.active if sheet_name not in wb.sheetnames else wb[sheet_name]
        header_row = [cell.value for cell in ws[1]]
        col_map = {str(val).strip(): idx + 1 for idx, val in enumerate(header_row) if val is not None}
        show_id_col_idx = next((idx for name, idx in col_map.items() if 'show id' in name.lower()), 1)
        
        excel_row_by_show_id = {}
        for r in range(2, ws.max_row + 1):
            raw_val = str(ws.cell(r, show_id_col_idx).value or '').strip()
            match = re.search(r'[a-f0-9]{40}', raw_val, re.I)
            clean_id = match.group(0).lower() if match else ''
            if clean_id:
                excel_row_by_show_id[clean_id] = r
                
        print(f"Matched {len(excel_row_by_show_id)} shows in '{ws.title}' sheet.")
        print("\n" + "="*60)
        print(" STARTING THE SCRAPER LOOP (ARGUS SHOW DATA ONLY)")
        print("="*60)
        
        total_shows = len(excel_row_by_show_id)
        for i, (show_id, r_idx) in enumerate(excel_row_by_show_id.items(), 1):
            url = f"https://cms.pocketfm.com/shows/audiobooks?tab=published&id={show_id}"
            print(f"\n[{i}/{total_shows}] Scraping Show ID: {show_id} (Row {r_idx})")
            
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=25000)
                
                # Check for redirect to login or logout
                page.wait_for_timeout(2000)
                body_check = page.inner_text("body").lower()
                if "login with jumpcloud" in body_check or "logout" in page.url.lower() or "jumpcloud.com" in page.url.lower():
                    print("\n[!] Session expired. Please log in on the browser window...")
                    page.goto("https://cms.pocketfm.com/")
                    input("Once logged in and the CMS dashboard renders, press ENTER here: ")
                    context.storage_state(path=auth_file)
                    page.goto(url, wait_until="domcontentloaded", timeout=25000)
                    page.wait_for_timeout(2000)

                # Wait for dashboard metrics to render (wait specifically for Plays stat card)
                try:
                    page.locator('div.sc-fmNzgT:has-text("Plays"), div:has-text("Plays")').first.wait_for(state="visible", timeout=15000)
                except Exception:
                    pass
                time.sleep(2) # Extra buffer for ratings and stats to finish fetching
                
                html = page.content()
                metrics = extract_metrics(html)
                
                for key in ['Plays', 'Duration', 'Ratings', 'Reviews', 'Comments']:
                    val = metrics.get(key, "").strip()
                    if not val or "no " in val.lower():
                        metrics[key] = "0"
                
                print(f"--> Found Name: {metrics['Show Name']}")
                print(f"--> Plays: {metrics['Plays']} | Duration: {metrics['Duration']} | Ratings: {metrics['Ratings']} | Reviews: {metrics['Reviews']} | Comments: {metrics['Comments']}")
                
                # Save immediately to Excel row
                for field in ['Show Name', 'Plays', 'Duration', 'Ratings', 'Reviews', 'Comments']:
                    if field in col_map:
                        val = metrics[field]
                        if field in ['Duration', 'Ratings']:
                            try:
                                ws.cell(r_idx, col_map[field]).value = float(val)
                            except Exception:
                                ws.cell(r_idx, col_map[field]).value = str(val)
                        elif field in ['Reviews', 'Comments']:
                            try:
                                ws.cell(r_idx, col_map[field]).value = float(val) if '.' in str(val) else int(val)
                            except Exception:
                                ws.cell(r_idx, col_map[field]).value = str(val)
                        else:
                            ws.cell(r_idx, col_map[field]).value = str(val)
                wb.save(excel_path)
                print(f"--> [SAVED] Successfully written to row {r_idx} in '{ws.title}'!")
                
            except Exception as e:
                print(f"--> Timeout or error loading {show_id}: {e}")
                
        print("\n" + "="*60)
        print(f"FINISHED SCRAPING! All metrics saved safely to '{excel_path}' -> [{ws.title}]!")
        print("="*60)
        browser.close()

if __name__ == "__main__":
    main()
