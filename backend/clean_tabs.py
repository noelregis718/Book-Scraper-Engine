import os
import openpyxl

folder = 'e:/Internship/PocketFM'
target_file = None

for file in os.listdir(folder):
    if 'episode metrics' in file.lower() and 'noel' in file.lower() and 'new' in file.lower() and file.endswith('.xlsx'):
        target_file = os.path.join(folder, file)
        break

if not target_file:
    print('Could not find the file.')
else:
    print(f'Found file: {target_file}')
    wb = openpyxl.load_workbook(target_file)
    sheets_deleted = 0
    
    target_sheet_name = None
    for name in wb.sheetnames:
        if name.lower().strip() == 'metrics 1':
            target_sheet_name = name
            break
            
    if not target_sheet_name:
        print('Could not find the tab "metrics 1" inside the file.')
        print(f'Available tabs are: {wb.sheetnames}')
    else:
        for sheet in wb.sheetnames:
            if sheet != target_sheet_name:
                del wb[sheet]
                sheets_deleted += 1
                
        wb.save(target_file)
        print(f'Successfully deleted {sheets_deleted} tabs, leaving only "{target_sheet_name}".')
