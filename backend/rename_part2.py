import openpyxl
import os
import re

file_path = r'e:\Internship\PocketFM\Romantasy - Subjective Reviews - 9.9.26.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Picks for Review (Sep 26)']

part2_rows = [192, 193, 196, 197, 203, 204, 211, 215, 216, 221, 223, 225, 227, 228, 235, 239, 247, 250, 251, 252, 256, 257, 258, 259, 262, 263, 266, 268, 269, 273, 275, 276, 289, 290, 297, 303, 306, 310, 312, 314, 315, 316, 318, 320, 322, 323, 325, 329, 330, 332, 333, 334, 335, 338, 339, 343, 349, 351, 353, 354, 357, 360, 362, 364, 365, 370, 371, 372, 375, 376, 378, 379, 380, 381, 382, 615, 616, 618, 619, 620, 622, 623, 628, 703, 704, 709, 710, 712, 716, 718, 721, 722, 723, 725, 727, 729, 730, 731, 860, 864, 866, 867, 868, 869, 871, 876, 877, 878, 880, 881, 882, 883, 885, 887, 891, 893, 894, 897, 899, 902, 906, 908, 909, 910, 961, 990, 1004, 1005, 1007, 1008]

part2_downloads = r'e:\Internship\PocketFM\downloads_part2'
folders = [d for d in os.listdir(part2_downloads) if os.path.isdir(os.path.join(part2_downloads, d))]

def clean_name(n):
    return re.sub(r'[^a-zA-Z0-9]', '', n).lower()

renamed = 0
for row_idx in part2_rows:
    series_name = str(ws.cell(row=row_idx, column=1).value or '').strip()
    author = str(ws.cell(row=row_idx, column=2).value or '').strip()
    
    target_clean = clean_name(author + series_name)
    target_series_only = clean_name(series_name)
    
    for folder in folders:
        # folder is something like "1002_Tournion"
        if '_' in folder:
            folder_series = folder.split('_', 1)[1]
        else:
            folder_series = folder
            
        folder_clean = clean_name(folder_series)
        
        if folder.startswith(str(row_idx) + '_'):
            continue
            
        if target_series_only == folder_clean:
            old_path = os.path.join(part2_downloads, folder)
            
            invalid_chars = ['<', '>', ':', '"', '/', '\\', '|', '?', '*']
            clean_series = series_name
            for char in invalid_chars:
                clean_series = clean_series.replace(char, '')
            clean_series = clean_series.strip()
            
            new_folder_name = f"{row_idx}_{clean_series}"
            new_path = os.path.join(part2_downloads, new_folder_name)
            
            if not os.path.exists(new_path):
                os.rename(old_path, new_path)
                renamed += 1
            break
            
print(f'Renamed {renamed} folders to have the correct new row IDs!')
