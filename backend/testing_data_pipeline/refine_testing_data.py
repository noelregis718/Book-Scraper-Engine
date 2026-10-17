import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_refined_sheets():
    src_file = 'Testing_Data_Genre_Wise.xlsx'
    refined_standalone_file = 'Testing_Data_Genre_Wise_Refined.xlsx'

    print(f"Loading {src_file}...")
    wb_in = openpyxl.load_workbook(src_file, data_only=True)

    rd_files = set(os.listdir('Downloaded_SRTs/RD_Drama'))
    ww_files = set(os.listdir('Downloaded_SRTs/Werewolf_Data'))

    print(f"Downloaded SRT count: RD_Drama={len(rd_files)}, Werewolf_Data={len(ww_files)}")

    # Column limit: 1 to 29 (A to AC)
    COL_LIMIT = 29

    # Styling definitions
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Deep slate navy
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='Segoe UI', size=10, color='0F172A')
    fill_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='E2E8F0')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style='medium', color='0F172A'))

    num_formats = {
        1: 'yyyy-mm-dd',
        11: '$#,##0.00',
        12: '$#,##0.000',
        13: '$#,##0.00',
        14: '0.0000%',
        15: '0.00%',
        16: '0.00%',
        17: '0.00%',
        18: '0.00%',
        19: '0.00',
        20: '0.00%',
        21: '0.00%',
        22: '0.00%',
        23: '0.00%',
        24: '0.00%',
        25: '#,##0',
        26: '#,##0',
        27: '#,##0',
        28: '0.00%'
    }

    alignments = {
        1: Alignment(horizontal='center', vertical='center'),
        7: Alignment(horizontal='center', vertical='center'),
        11: Alignment(horizontal='right', vertical='center'),
        12: Alignment(horizontal='right', vertical='center'),
        13: Alignment(horizontal='right', vertical='center'),
        14: Alignment(horizontal='right', vertical='center'),
        15: Alignment(horizontal='right', vertical='center'),
        16: Alignment(horizontal='right', vertical='center'),
        17: Alignment(horizontal='right', vertical='center'),
        18: Alignment(horizontal='right', vertical='center'),
        19: Alignment(horizontal='right', vertical='center'),
        20: Alignment(horizontal='right', vertical='center'),
        21: Alignment(horizontal='right', vertical='center'),
        22: Alignment(horizontal='right', vertical='center'),
        23: Alignment(horizontal='right', vertical='center'),
        24: Alignment(horizontal='right', vertical='center'),
        25: Alignment(horizontal='right', vertical='center'),
        26: Alignment(horizontal='right', vertical='center'),
        27: Alignment(horizontal='right', vertical='center'),
        28: Alignment(horizontal='right', vertical='center'),
        29: Alignment(horizontal='center', vertical='center'), # Asset Code
    }
    default_alignment = Alignment(horizontal='left', vertical='center')

    # Extract rows for RD_Drama
    ws_rd_in = wb_in['RD_Drama']
    rd_headers = [ws_rd_in.cell(row=1, column=c).value for c in range(1, COL_LIMIT + 1)]
    rd_rows = []
    seen_rd = set()
    for r in range(2, ws_rd_in.max_row + 1):
        asset = str(ws_rd_in.cell(row=r, column=29).value or '').strip()
        if f"{asset}.srt" in rd_files and asset not in seen_rd:
            seen_rd.add(asset)
            row_data = [ws_rd_in.cell(row=r, column=c).value for c in range(1, COL_LIMIT + 1)]
            rd_rows.append(row_data)

    # Extract rows for Werewolf_Data
    ws_ww_in = wb_in['Werewolf_Data']
    ww_headers = [ws_ww_in.cell(row=1, column=c).value for c in range(1, COL_LIMIT + 1)]
    ww_rows = []
    seen_ww = set()
    for r in range(2, ws_ww_in.max_row + 1):
        asset = str(ws_ww_in.cell(row=r, column=29).value or '').strip()
        if f"{asset}.srt" in ww_files and asset not in seen_ww:
            seen_ww.add(asset)
            row_data = [ws_ww_in.cell(row=r, column=c).value for c in range(1, COL_LIMIT + 1)]
            ww_rows.append(row_data)

    print(f"Extracted: RD_Drama = {len(rd_rows)} rows, Werewolf_Data = {len(ww_rows)} rows.")

    def format_sheet(ws, headers, rows):
        ws.append(headers)
        ws.row_dimensions[1].height = 32

        for c_idx in range(1, COL_LIMIT + 1):
            cell = ws.cell(row=1, column=c_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = header_border

        for r_idx, r_data in enumerate(rows, start=2):
            ws.append(r_data)
            ws.row_dimensions[r_idx].height = 20
            row_fill = fill_even if r_idx % 2 == 0 else fill_odd

            for c_idx in range(1, COL_LIMIT + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.fill = row_fill
                cell.border = cell_border
                cell.alignment = alignments.get(c_idx, default_alignment)

                if c_idx in num_formats and cell.value is not None:
                    cell.number_format = num_formats[c_idx]

        # Calculate width
        for c_idx in range(1, COL_LIMIT + 1):
            col_letter = get_column_letter(c_idx)
            h_text = str(headers[c_idx - 1] or '')
            max_len = len(h_text)
            for r_data in rows[:150]:
                val = r_data[c_idx - 1]
                if val is not None:
                    max_len = max(max_len, len(str(val)))

            if c_idx == 1:
                w = 14
            elif c_idx == 7:
                w = 38
            elif c_idx in [2, 3, 4]:
                w = min(max(max_len + 3, 20), 45)
            elif c_idx in [5, 6]:
                w = min(max(max_len + 3, 16), 35)
            elif c_idx in [11, 12, 13]:
                w = 16
            elif c_idx == 29:
                w = 16
            else:
                w = min(max(max_len + 3, 12), 26)

            ws.column_dimensions[col_letter].width = w

        ws.freeze_panes = 'A2'
        last_col = get_column_letter(COL_LIMIT)
        total_rows = len(rows) + 1
        ws.auto_filter.ref = f"A1:{last_col}{total_rows}"
        ws.views.sheetView[0].showGridLines = True
        ws.protection.sheet = False

    # 1. Update Testing_Data_Genre_Wise.xlsx with refined tabs as primary tabs
    # We load Testing_Data_Genre_Wise.xlsx without data_only to avoid losing anything, or we build a fresh workbook containing both refined and existing
    wb_full = openpyxl.load_workbook(src_file)

    # Remove existing refined tabs if already present
    for t in ['RD_Drama_Refined', 'Werewolf_Data_Refined']:
        if t in wb_full.sheetnames:
            del wb_full[t]

    # Create RD_Drama_Refined and Werewolf_Data_Refined at the front
    ws_rd_ref = wb_full.create_sheet(title='RD_Drama_Refined', index=0)
    format_sheet(ws_rd_ref, rd_headers, rd_rows)

    ws_ww_ref = wb_full.create_sheet(title='Werewolf_Data_Refined', index=1)
    format_sheet(ws_ww_ref, ww_headers, ww_rows)

    wb_full.save(src_file)
    wb_full.close()
    print(f"Saved refined tabs into {src_file}.")

    # 2. Also create standalone Testing_Data_Genre_Wise_Refined.xlsx
    wb_stand = openpyxl.Workbook()
    wb_stand.remove(wb_stand.active)

    ws_s_rd = wb_stand.create_sheet(title='RD_Drama')
    format_sheet(ws_s_rd, rd_headers, rd_rows)

    ws_s_ww = wb_stand.create_sheet(title='Werewolf_Data')
    format_sheet(ws_s_ww, ww_headers, ww_rows)

    wb_stand.save(refined_standalone_file)
    wb_stand.close()
    print(f"Saved standalone refined workbook to {refined_standalone_file}.")

if __name__ == '__main__':
    build_refined_sheets()
