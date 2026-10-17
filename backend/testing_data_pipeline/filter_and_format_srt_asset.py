import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def process_file():
    src_file = 'Testing_Data_Genre_Wise_ORIGINAL_BACKUP.xlsx'
    out_file = 'Testing_Data_Genre_Wise.xlsx'

    print(f"Loading {src_file} with data_only=True...")
    wb_in = openpyxl.load_workbook(src_file, data_only=True, read_only=True)
    wb_out = openpyxl.Workbook()
    default_sheet = wb_out.active

    # Styling definitions
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid') # Deep slate navy
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='Segoe UI', size=10, color='0F172A')
    fill_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='E2E8F0')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style='medium', color='0F172A'))

    num_formats = {
        1: 'yyyy-mm-dd',
        11: '$#,##0.00',
        12: '$#,##0.000',
        13: '$#,##0.00',
        14: '0.0000%',
        15: '0.00%',
        16: '0.00%',
        17: '0.00%',
        18: '0.00%',
        19: '0.00',
        20: '0.00%',
        21: '0.00%',
        22: '0.00%',
        23: '0.00%',
        24: '0.00%',
        25: '#,##0',
        26: '#,##0',
        27: '#,##0',
        28: '0.00%'
    }

    alignments = {
        1: Alignment(horizontal='center', vertical='center'),
        7: Alignment(horizontal='center', vertical='center'),
        11: Alignment(horizontal='right', vertical='center'),
        12: Alignment(horizontal='right', vertical='center'),
        13: Alignment(horizontal='right', vertical='center'),
        14: Alignment(horizontal='right', vertical='center'),
        15: Alignment(horizontal='right', vertical='center'),
        16: Alignment(horizontal='right', vertical='center'),
        17: Alignment(horizontal='right', vertical='center'),
        18: Alignment(horizontal='right', vertical='center'),
        19: Alignment(horizontal='right', vertical='center'),
        20: Alignment(horizontal='right', vertical='center'),
        21: Alignment(horizontal='right', vertical='center'),
        22: Alignment(horizontal='right', vertical='center'),
        23: Alignment(horizontal='right', vertical='center'),
        24: Alignment(horizontal='right', vertical='center'),
        25: Alignment(horizontal='right', vertical='center'),
        26: Alignment(horizontal='right', vertical='center'),
        27: Alignment(horizontal='right', vertical='center'),
        28: Alignment(horizontal='right', vertical='center'),
        29: Alignment(horizontal='center', vertical='center'), # Asset Code
    }
    default_alignment = Alignment(horizontal='left', vertical='center')

    target_sheets = [
        ('RD_Drama', 37),      # 37 columns
        ('Werewolf_Data', 38)  # 38 columns
    ]

    for sname, col_count in target_sheets:
        print(f"\nProcessing and filtering: {sname} (Total columns: {col_count})...")
        ws_in = wb_in[sname]
        ws_out = wb_out.create_sheet(title=sname)

        # Read header row
        header_row = None
        filtered_rows = []

        for r_idx, row in enumerate(ws_in.iter_rows(values_only=True)):
            if r_idx == 0:
                header_row = list(row[:col_count])
                continue

            if not any(row):
                continue

            asset = row[28] if len(row) > 28 else None
            srt = row[34] if len(row) > 34 else None

            # Filter out blanks in both Asset Code (Col 29) and SRT (Col 35)
            is_asset_valid = asset is not None and str(asset).strip() not in ['', 'None', 'null']
            is_srt_valid = srt is not None and str(srt).strip() not in ['', 'None', 'null']

            if is_asset_valid and is_srt_valid:
                filtered_rows.append(list(row[:col_count]))

        print(f"  {sname}: Filtered down to {len(filtered_rows)} rows where both Asset Code and SRT are non-blank.")

        # Write Header
        ws_out.append(header_row)
        ws_out.row_dimensions[1].height = 32
        for col_idx in range(1, col_count + 1):
            c = ws_out.cell(row=1, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_alignment
            c.border = header_border

        # Write Data Rows & apply formatting
        for row_idx, r_data in enumerate(filtered_rows, start=2):
            ws_out.append(r_data)
            ws_out.row_dimensions[row_idx].height = 20
            row_fill = fill_even if row_idx % 2 == 0 else fill_odd

            for col_idx in range(1, col_count + 1):
                cell = ws_out.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.fill = row_fill
                cell.border = cell_border
                cell.alignment = alignments.get(col_idx, default_alignment)

                if col_idx in num_formats and cell.value is not None:
                    cell.number_format = num_formats[col_idx]

        # Calculate Column Widths
        print(f"  Calculating tailored column widths for {sname}...")
        for col_idx in range(1, col_count + 1):
            col_letter = get_column_letter(col_idx)
            h_text = str(header_row[col_idx - 1] or '')
            max_len = len(h_text)

            for sample_row in filtered_rows[:100]:
                val = sample_row[col_idx - 1] if col_idx - 1 < len(sample_row) else None
                if val is not None:
                    max_len = max(max_len, len(str(val)))

            if col_idx == 1:
                width = 14
            elif col_idx == 7:
                width = 38
            elif col_idx in [2, 3, 4]:
                width = min(max(max_len + 3, 20), 45)
            elif col_idx in [5, 6]:
                width = min(max(max_len + 3, 16), 35)
            elif col_idx in [11, 12, 13]:
                width = 16
            elif col_idx == 29:
                width = 16 # Asset Code
            elif col_idx == 35:
                width = 38 # SRT
            elif col_idx in [30, 31, 32, 33, 34]:
                width = 25
            else:
                width = min(max(max_len + 3, 12), 28)

            ws_out.column_dimensions[col_letter].width = width

        # View settings: Freeze Panes & AutoFilter
        ws_out.freeze_panes = 'A2'
        last_col_letter = get_column_letter(col_count)
        total_out_rows = len(filtered_rows) + 1
        ws_out.auto_filter.ref = f"A1:{last_col_letter}{total_out_rows}"
        ws_out.views.sheetView[0].showGridLines = True
        ws_out.protection.sheet = False # 100% editable

    wb_out.remove(default_sheet)
    wb_in.close()

    print(f"\nSaving final filtered workbook to {out_file}...")
    wb_out.save(out_file)
    wb_out.close()
    print("Done! Both tabs successfully filtered and professionally styled.")

if __name__ == '__main__':
    process_file()
