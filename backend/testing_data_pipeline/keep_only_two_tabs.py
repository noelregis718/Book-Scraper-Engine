import os
import shutil
import openpyxl

def main():
    src_file = 'Testing_Data_Genre_Wise.xlsx'
    backup_file = 'Testing_Data_Genre_Wise_ORIGINAL_BACKUP.xlsx'
    temp_output_file = 'Testing_Data_Genre_Wise_NEW.xlsx'

    # Step 1: Backup original file if not already backed up
    if not os.path.exists(backup_file):
        print(f"Creating safety backup copy: {backup_file}...")
        shutil.copy2(src_file, backup_file)
        print("Safety backup created successfully.")
    else:
        print(f"Backup {backup_file} already exists.")

    # Step 2: Read only RD_Drama and Werewolf_Data with data_only=True
    print("Loading source workbook (data_only=True, read_only=True)...")
    wb_in = openpyxl.load_workbook(src_file, data_only=True, read_only=True)

    # Step 3: Create a clean new workbook with only RD_Drama and Werewolf_Data
    wb_out = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb_out.active

    target_sheets = ['RD_Drama', 'Werewolf_Data']

    for sname in target_sheets:
        print(f"Processing sheet: {sname}...")
        ws_in = wb_in[sname]
        ws_out = wb_out.create_sheet(title=sname)

        row_count = 0
        for row in ws_in.iter_rows(values_only=True):
            # Only append if row has at least one non-None value
            if any(cell is not None for cell in row):
                ws_out.append(list(row[:34]))
                row_count += 1
                if row_count % 5000 == 0:
                    print(f"  Appended {row_count} rows in {sname}...")
        print(f"Finished {sname}: Total {row_count} rows written.")

    # Remove initial default empty sheet
    wb_out.remove(default_sheet)

    # Save to temp output file
    print(f"Saving trimmed workbook to {temp_output_file}...")
    wb_out.save(temp_output_file)
    wb_in.close()
    wb_out.close()
    print("Save complete.")

    # Step 4: Replace Testing_Data_Genre_Wise.xlsx with the trimmed version
    print(f"Replacing {src_file} with trimmed file...")
    shutil.move(temp_output_file, src_file)
    print("Done! Only RD_Drama and Werewolf_Data remain in Testing_Data_Genre_Wise.xlsx.")

if __name__ == '__main__':
    main()
