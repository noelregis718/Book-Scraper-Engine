import openpyxl
import os

wb = openpyxl.load_workbook('Testing_Data_Genre_Wise.xlsx', data_only=True)

for sname in ['RD_Drama', 'Werewolf_Data']:
    ws = wb[sname]
    max_r = ws.max_row
    if sname == 'RD_Drama':
        target_rows = [17, 50, 150, 450, 750]
    else:
        target_rows = [2, 35, 120, 300, 550]

    print(f"\n=======================================================")
    print(f"=== {sname} (Total Data Rows: {max_r - 1}) ===")
    print(f"=======================================================")
    for r in target_rows:
        show = ws.cell(row=r, column=6).value
        asset = str(ws.cell(row=r, column=29).value or '').strip()
        srt_url = str(ws.cell(row=r, column=35).value or '').strip()

        file_name = f"{asset}.srt"
        file_path = os.path.join('Downloaded_SRTs', sname, file_name)
        exists = os.path.exists(file_path)

        first_dialogue = ""
        if exists:
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as fp:
                    content = fp.read()
                    lines = [l.strip() for l in content.splitlines() if l.strip() and not l.strip().isdigit() and '-->' not in l]
                    first_dialogue = " ".join(lines[:2])
            except Exception as e:
                first_dialogue = str(e)

        print(f"• Row {r}:")
        print(f"    Show Name    : {show}")
        print(f"    Asset Code   : {asset}")
        print(f"    Saved File   : {file_name} (in Downloaded_SRTs/{sname}/)")
        print(f"    First Dialogue: \"{first_dialogue[:90]}...\"")
