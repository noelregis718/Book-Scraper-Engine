import openpyxl
import os
import shutil
import subprocess
import re

file_path = r"e:\Internship\PocketFM\RB Media _ Pocket FM List - 2 (Testing Final) _ Internal Copy.xlsx"
output_dir = r"e:\Internship\PocketFM\Drive_Folder_Zips"

def sanitize_filename(name):
    # Remove invalid characters for Windows filenames
    return re.sub(r'[\\/*?:"<>|]', "", str(name)).strip()

def main():
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    print(f"Loading {file_path}...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb["Testing List metrics"]
    
    # Extract links from column N (14)
    links = []
    for row in range(2, ws.max_row + 1):
        # Only process rows where the Genre (Column C / 3) is Romantasy
        genre = str(ws.cell(row=row, column=3).value).strip().lower()
        if "romantasy" not in genre:
            continue
            
        cell = ws.cell(row=row, column=14)
        if cell.hyperlink and "drive.google.com" in str(cell.hyperlink.target):
            # Try to get series name from cell text, fallback to Column B (Series Name)
            name = str(cell.value).strip()
            if not name or name.lower() == "link" or name.lower() == "none":
                name = str(ws.cell(row=row, column=2).value).strip()
                
            links.append((name, cell.hyperlink.target))
            
    print(f"Found {len(links)} Google Drive folders to download.")
    
    for idx, (name, url) in enumerate(links, 1):
        safe_name = sanitize_filename(name)
        if not safe_name or safe_name.lower() == "none":
            safe_name = f"Unknown_Folder_{idx}"
            
        zip_path = os.path.join(output_dir, f"{safe_name}.zip")
        
        # Skip if already exists
        if os.path.exists(zip_path):
            print(f"\n[{idx}/{len(links)}] Skipping '{safe_name}' - Already downloaded and zipped.")
            continue
            
        print(f"\n[{idx}/{len(links)}] Downloading '{safe_name}'...")
        
        # Create a temp download dir specific to this book
        temp_dl_dir = os.path.join(output_dir, f"temp_{idx}")
        if not os.path.exists(temp_dl_dir):
            os.makedirs(temp_dl_dir)
            
        try:
            # Run gdown. It will output directly to the terminal so user sees progress.
            subprocess.run(
                ["python", "-m", "gdown", "--folder", url],
                cwd=temp_dl_dir,
            )
            
            # Check if it actually downloaded anything
            items = os.listdir(temp_dl_dir)
            if not items:
                print(f"   -> Error: Failed to download or folder is completely empty.")
                continue
                
            print(f"   -> Download complete. Compressing to .zip...")
            # Zips the contents of the temp directory directly
            shutil.make_archive(zip_path.replace('.zip', ''), 'zip', temp_dl_dir)
            
            print(f"   -> Success! Saved to {safe_name}.zip")
        except Exception as e:
            print(f"   -> Error processing {safe_name}: {e}")
        finally:
            # Cleanup raw unzipped files immediately
            if os.path.exists(temp_dl_dir):
                shutil.rmtree(temp_dl_dir, ignore_errors=True)

    print("\nALL DONE! All Google Drive folders have been downloaded and compressed into ZIP files.")

if __name__ == "__main__":
    main()
